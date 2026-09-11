import openpyxl, datetime as dt
from collections import defaultdict

wbv = openpyxl.load_workbook('ai_referral_ctr.xlsx', data_only=True)
wbf = openpyxl.load_workbook('ai_referral_ctr.xlsx')

def rows(ws, hdr=1):
    heads=[c.value for c in ws[hdr]]
    out=[]
    for r in ws.iter_rows(min_row=hdr+1, values_only=True):
        out.append(dict(zip(heads,r)))
    return out

S = rows(wbv['Sessions'])
S = [r for r in S if isinstance(r['date'], dt.datetime)]
P = [r for r in rows(wbv['Pages']) if isinstance(r['canonical_url'],str) and r['canonical_url'].startswith('http')]
D = [r for r in rows(wbv['Daily']) if isinstance(r['date'], dt.datetime)]
W = [r for r in rows(wbv['Weekly']) if isinstance(r['week starting'], dt.datetime)]
R = [r for r in rows(wbv['Referrers']) if r['referrer'] in ('ChatGPT','Claude','Gemini','Copilot')]

print(f"Sessions data rows: {len(S)}  Pages: {len(P)}  Daily: {len(D)}  Weekly: {len(W)}  Referrers: {len(R)}")
issues=[]
def chk(name, got, exp, tol=1e-9):
    ok = (abs(got-exp)<=tol) if isinstance(exp,(int,float)) and isinstance(got,(int,float)) else got==exp
    if not ok: issues.append(f"{name}: sheet={got!r} recomputed={exp!r}")
    return ok

# ---------- 1. Sessions integrity ----------
tot_sessions = sum(r['sessions'] for r in S)
print("\n-- Sessions --")
print("total sessions:", tot_sessions)
print("sessions values distinct:", sorted(set(r['sessions'] for r in S)))
print("date range:", min(r['date'] for r in S).date(), "->", max(r['date'] for r in S).date())
print("referrers:", sorted(set(r['referrer'] for r in S)))
# duplicate grain check: date+referrer+url should be unique
g=defaultdict(int)
for r in S: g[(r['date'],r['referrer'],r['canonical_url'])]+=1
dups={k:v for k,v in g.items() if v>1}
print("duplicate (date,referrer,url) rows:", len(dups))
for k,v in list(dups.items())[:5]: print("   ",k[0].date(),k[1],k[2][-60:],v)
# verify 'first row for page+date' flag
seen=set(); exp_flag=[]
for r in S:
    key=(r['canonical_url'], r['date'])
    exp_flag.append(0 if key in seen else 1); seen.add(key)
bad=[i for i,(r,e) in enumerate(zip(S,exp_flag)) if r['first row for page+date']!=e]
print("first-row-flag mismatches:", len(bad))

# ---------- 2. Pages tab ----------
print("\n-- Pages --")
urls=[r['canonical_url'] for r in P]
print("duplicate canonical_url in Pages:", len(urls)-len(set(urls)))
sess_urls=set(r['canonical_url'] for r in S)
print("session URLs missing from Pages:", len(sess_urls-set(urls)))
for u in list(sess_urls-set(urls))[:5]: print("   ", u)

by_url=defaultdict(lambda: defaultdict(int))
for r in S:
    by_url[r['canonical_url']]['all'] += r['sessions']
    by_url[r['canonical_url']][r['referrer']] += r['sessions']
    by_url[r['canonical_url']]['days'] += r['first row for page+date']

E_tot = sum(r['citations (Bing AI)'] or 0 for r in P)
F_tot = sum(r['sessions (all AI referrers)'] or 0 for r in P)
chk("Pages E106 citations total", wbv['Pages']['E106'].value, E_tot)
chk("Pages F106 sessions total", wbv['Pages']['F106'].value, F_tot)
chk("Pages F106 vs Sessions tab total", wbv['Pages']['F106'].value, tot_sessions)
for col,ref in [('sessions (all AI referrers)','all'),('ChatGPT','ChatGPT'),('Claude','Claude'),('Gemini','Gemini'),('Copilot','Copilot'),('days with sessions','days')]:
    n=0
    for r in P:
        exp = by_url[r['canonical_url']][ref]
        if (r[col] or 0) != exp: n+=1
    if n: issues.append(f"Pages col '{col}': {n} rows disagree with Sessions recompute")
print("per-page SUMIFS recompute mismatches:", "none" if not any('Pages col' in i for i in issues) else [i for i in issues if 'Pages col' in i])

# CTR / share columns
ctr_bad=share_bad=0
for r in P:
    e=r['citations (Bing AI)']; f=r['sessions (all AI referrers)'] or 0; j=r['Copilot'] or 0
    exp = '' if not e else f/e
    got = r['CTR all referrers']
    got = '' if got is None else got
    if isinstance(exp,float) and (not isinstance(got,float) or abs(got-exp)>1e-9): ctr_bad+=1
    if exp=='' and isinstance(got,float): ctr_bad+=1
    m = r['share of citations']; mexp=(e or 0)/E_tot
    n_ = r['share of sessions']; nexp=f/F_tot
    if abs((m or 0)-mexp)>1e-9 or abs((n_ or 0)-nexp)>1e-9: share_bad+=1
print("page CTR mismatches:", ctr_bad, " share mismatches:", share_bad)
print("sum share of citations:", round(sum(r['share of citations'] or 0 for r in P),10))
print("sum share of sessions:", round(sum(r['share of sessions'] or 0 for r in P),10))

# match status consistency
ms=defaultdict(int)
ms_bad=[]
for r in P:
    e=r['citations (Bing AI)']; f=r['sessions (all AI referrers)'] or 0
    st=r['match status']; ms[st]+=1
    if e and f: exp='cited and landed'
    elif e and not f: exp='cited, no landings'
    elif (not e) and f: exp='landings, not cited'
    else: exp='neither'
    if st!=exp: ms_bad.append((r['path'],st,exp,e,f))
print("match status counts:", dict(ms))
print("match status mismatches:", len(ms_bad))
for x in ms_bad[:10]: print("   ", x)

# README counts
cited = sum(1 for r in P if str(r['match status']).startswith('cited'))
landed = sum(1 for r in P if str(r['match status']).endswith('landed')) + sum(1 for r in P if str(r['match status']).startswith('landings'))
print("README 'Pages cited by Bing AI' COUNTIF recompute:", cited)
print("README 'Pages with AI referral landings' recompute:", landed)
print("  (pages with >0 sessions, truth):", sum(1 for r in P if (r['sessions (all AI referrers)'] or 0)>0))
print("  (pages with >0 citations, truth):", sum(1 for r in P if (r['citations (Bing AI)'] or 0)>0))

# ---------- 3. Daily tab ----------
print("\n-- Daily --")
by_date=defaultdict(lambda: defaultdict(int))
for r in S:
    by_date[r['date']]['all']+=r['sessions']
    by_date[r['date']][r['referrer']]+=r['sessions']
dmis=0
for r in D:
    for col,ref in [('sessions (all AI referrers)','all'),('ChatGPT','ChatGPT'),('Claude','Claude'),('Gemini','Gemini'),('Copilot','Copilot')]:
        if (r[col] or 0)!=by_date[r['date']][ref]: dmis+=1
print("Daily SUMIFS mismatches:", dmis)
dates=[r['date'] for r in D]
print("Daily span:", min(dates).date(),'->',max(dates).date(), "rows:", len(D), "expected contiguous:", (max(dates)-min(dates)).days+1)
print("gaps:", [ (min(dates)+dt.timedelta(days=i)).date() for i in range((max(dates)-min(dates)).days+1) if (min(dates)+dt.timedelta(days=i)) not in set(dates)])
cit_days=[r for r in D if r['citations (Bing AI)'] is not None]
sess_days=[r for r in D if (r['sessions (all AI referrers)'] or 0)>0]
print("citation coverage:", min(r['date'] for r in cit_days).date(),'->',max(r['date'] for r in cit_days).date(), len(cit_days),'days')
print("session dates:", min(r['date'] for r in S).date(),'->',max(r['date'] for r in S).date())
H=sum(r['citations (Bing AI)'] or 0 for r in D)
chk("Daily H123 citations total", wbv['Daily']['H123'].value, H)
chk("Daily J123 sessions total", wbv['Daily']['J123'].value, sum(r['sessions (all AI referrers)'] or 0 for r in D))
avg_cited_pages = sum(r['cited pages'] for r in cit_days)/len(cit_days)
chk("Daily I123 avg cited pages", wbv['Daily']['I123'].value, avg_cited_pages, 1e-9)
# coverage flag
cov_bad=[]
sess_dates=set(r['date'] for r in S)
for r in D:
    has_c = r['citations (Bing AI)'] is not None
    has_s = r['date'] in sess_dates
    exp = 'both' if has_c and has_s else ('citations only' if has_c else ('sessions only' if has_s else 'neither'))
    if r['source coverage']!=exp: cov_bad.append((r['date'].date(), r['source coverage'], exp))
print("source-coverage flag mismatches:", len(cov_bad), cov_bad[:10])
# ISO week / weekday / weekend checks
wk_bad=[r['date'].date() for r in D if r['ISO week']!=r['date'].isocalendar()[1]]
wd_bad=[r['date'].date() for r in D if r['day']!=r['date'].strftime('%a')]
we_bad=[r['date'].date() for r in D if r['weekend']!=('yes' if r['date'].weekday()>=5 else 'no')]
ws_bad=[r['date'].date() for r in D if r['week starting']!=r['date']-dt.timedelta(days=r['date'].weekday())]
print("ISO week bad:",len(wk_bad),"weekday bad:",len(wd_bad),"weekend bad:",len(we_bad),"week-start bad:",len(ws_bad))

# ---------- 4. Weekly ----------
print("\n-- Weekly --")
wmis=0
for r in W:
    ws_=r['week starting']
    days=[d for d in D if d['week starting']==ws_]
    e=sum(d['citations (Bing AI)'] or 0 for d in days)
    f=sum(d['sessions (all AI referrers)'] or 0 for d in days)
    if (r['citations (Bing AI)'] or 0)!=e or (r['sessions (all AI referrers)'] or 0)!=f: wmis+=1; print("  wk",ws_.date(),r['citations (Bing AI)'],e,r['sessions (all AI referrers)'],f)
    if r['days in range']!=len(days): print("  days-in-range",ws_.date(),r['days in range'],len(days))
    dctr=[d['CTR all referrers'] for d in days if isinstance(d['CTR all referrers'],(int,float))]
    exp_l=sum(dctr)/len(dctr) if dctr else None
    if exp_l is not None and abs((r['CTR (mean of daily CTRs)'] or 0)-exp_l)>1e-9: print("  L mismatch",ws_.date(),r['CTR (mean of daily CTRs)'],exp_l)
print("Weekly rollup mismatches:", wmis)
chk("Weekly E20 vs Daily H123", wbv['Weekly']['E20'].value, wbv['Daily']['H123'].value)
chk("Weekly F20 vs Daily J123", wbv['Weekly']['F20'].value, wbv['Daily']['J123'].value)
print("Weekly L20 (avg of weekly avgs):", wbv['Weekly']['L20'].value, "| unweighted mean of ALL daily CTRs:",
      sum(d['CTR all referrers'] for d in D if isinstance(d['CTR all referrers'],(int,float)))/len([d for d in D if isinstance(d['CTR all referrers'],(int,float))]))

# ---------- 5. Referrers ----------
print("\n-- Referrers --")
for r in R:
    exp=sum(x['sessions'] for x in S if x['referrer']==r['referrer'])
    dp=len(set(x['canonical_url'] for x in S if x['referrer']==r['referrer']))
    fs=min(x['date'] for x in S if x['referrer']==r['referrer']).date()
    ls=max(x['date'] for x in S if x['referrer']==r['referrer']).date()
    print(f"  {r['referrer']:8} sheet={r['sessions']:4} recomputed={exp:4}  distinct pages sheet={r['distinct landing pages']} recomputed={dp}  first {r['first session'].date()}/{fs}  last {r['last session'].date()}/{ls}")
print("  TOTAL distinct pages sheet=", wbv['Referrers']['D6'].value, "recomputed=", len(set(x['canonical_url'] for x in S)))

# ---------- 6. README headline ----------
print("\n-- README headline --")
print("  citations(by page) E106 =", wbv['Pages']['E106'].value)
print("  citations(by day)  H123 =", wbv['Daily']['H123'].value)
d=wbv['Daily']['H123'].value; p=wbv['Pages']['E106'].value
print(f"  discrepancy: {d-p} ({(d-p)/d:.2%} of daily total, {(d-p)/p:.2%} of page total)")
print("  overall CTR (all AI):", wbv['Pages']['K106'].value)
print("  Copilot-only CTR:", wbv['Pages']['L106'].value, " = 3 /", p)
print("  CTR on daily denom:", 443/d)

# both-coverage restricted
both=[r for r in D if r['source coverage']=='both']
bc=sum(r['citations (Bing AI)'] for r in both); bs=sum(r['sessions (all AI referrers)'] or 0 for r in both)
print(f"  RESTRICTED to 'both' days ({len(both)} days): citations={bc} sessions={bs} CTR={bs/bc:.4%}")
print("  sessions on 'sessions only' days:", sum(r['sessions (all AI referrers)'] or 0 for r in D if r['source coverage']=='sessions only'))
print("  citations on 'citations only' days:", sum(r['citations (Bing AI)'] or 0 for r in D if r['source coverage']=='citations only'))

# outlier
top=sorted([r for r in cit_days], key=lambda r:-r['citations (Bing AI)'])[:6]
med=sorted(r['citations (Bing AI)'] for r in cit_days)[len(cit_days)//2]
print("\n-- Outlier --")
print("  median daily citations:", med)
for r in top: print(f"   {r['date'].date()} {r['citations (Bing AI)']:5}  = {r['citations (Bing AI)']/med:.1f}x median")
ex=[r for r in cit_days if r['date']!=dt.datetime(2026,7,30)]
print(f"  all-period CTR excl 30 Jul: {443/sum(r['citations (Bing AI)'] for r in ex):.4%} vs incl {443/d:.4%}")

# tracking params / 404
print("\n-- README claims --")
tp=sum(1 for r in S if r['had_tracking_params']=='yes')
tp_cg=sum(1 for r in S if r['had_tracking_params']=='yes' and r['referrer']=='ChatGPT')
cg=sum(1 for r in S if r['referrer']=='ChatGPT')
print(f"  rows with tracking params: {tp}/{len(S)} ({tp/len(S):.1%}); ChatGPT {tp_cg}/{cg} ({tp_cg/cg:.1%})")
print("  rows whose raw != canonical:", sum(1 for r in S if r['raw_landing_url']!=r['canonical_url']))
f404=[r for r in S if '404' in str(r['page_title']).lower() or '404' in str(r['canonical_url']).lower()]
print("  404-ish rows:", len(f404))
for r in f404: print("    ", r['date'].date(), r['referrer'], r['page_title'], '|', r['raw_landing_url'][:110])

print("\n================ ISSUES ================")
for i in issues: print(" *", i)
if not issues: print(" none from exact-match checks")
